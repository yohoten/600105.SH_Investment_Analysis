# -*- coding: utf-8 -*-
"""06_build_report.py —— 按《证券投资学》实践课报告模板与格式规范生成 Word 报告

格式规范（来自任务书）：
  页面 A4 / 上下 2.5cm / 左右 2cm / 页眉 1.5cm / 页脚 1.8cm
  页眉：宋体小五居中「XX企业投资分析报告」；页脚：居中页码
  一级标题 黑体三号加粗居中，段后 1.5 行
  二级标题 黑体四号加粗左对齐，段后 1 行
  三级标题 黑体小四加粗左对齐，段后 0.5 行
  正文 宋体小四，1.5 倍行距，首行缩进 2 字符
  摘要/关键词/参考文献 宋体小五
  图标题在图下方居中（宋体小五加粗）；表标题在表上方居中（宋体小五加粗）
  图表下方数据来源（宋体六号）
"""
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figures")
TPL = os.path.join(ROOT, "经济与管理学院实践课报告-《证券投资学》.docx")
OUT = os.path.join(ROOT, "12607320270457_滕宇豪_25级大数据与财务管理03班_江苏永鼎股份（600105.SH）投资分析报告.docx")

HEADER_TEXT = "永鼎股份（600105.SH）投资分析报告"


# ------------------------------------------------------------------ 基础工具
def set_font(run, cn="宋体", size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _fmt(p, line=1.5, before=0, after=0, indent=0, align=None, size=12):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(size * indent) if indent else Pt(0)
    if align is not None:
        p.alignment = align


def add_field(par, code="PAGE"):
    run = par.add_run()
    set_font(run, "宋体", 9)
    fld_b = OxmlElement("w:fldChar"); fld_b.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" {code} "
    fld_e = OxmlElement("w:fldChar"); fld_e.set(qn("w:fldCharType"), "end")
    run._r.append(fld_b)
    run._r.append(instr)
    run._r.append(fld_e)


def set_borders(table, sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


# tblPr 子元素的 schema 顺序（乱序会导致 Word 报“内容有问题”）
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
                "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]


def _order_tblPr(tblPr):
    idx = {qn(f"w:{n}"): i for i, n in enumerate(_TBLPR_ORDER)}
    kids = list(tblPr)
    kids.sort(key=lambda e: idx.get(e.tag, 99))
    for e in kids:
        tblPr.remove(e)
    for e in kids:
        tblPr.append(e)


def fix_table_width(table, widths, cell_mar=40):
    """固定列宽 + 固定布局，避免 Word 自动扩宽后表格越出版心。"""
    tblPr = table._tbl.tblPr
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))   # 1cm = 567 dxa
    tblPr.append(tblW)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:type"), "dxa")
        el.set(qn("w:w"), str(cell_mar))
        mar.append(el)
    tblPr.append(mar)
    _order_tblPr(tblPr)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def cell_text(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn="宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    _fmt(p, line=1.0, before=2, after=2, align=align, size=size)
    run = p.add_run(str(text))
    set_font(run, cn, size, bold)


# ------------------------------------------------------------------ 文档骨架
class Report:
    def __init__(self):
        self.doc = Document()
        self._setup_base_style()

    def _setup_base_style(self):
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)

    # ---------- 节与页眉页脚
    def config_section(self, sec, header=True, page_start=None, page_no=False):
        """配置节的页面、页眉、页脚。

        注意：python-docx 的 add_section() 会深拷贝上一节的 sectPr，
        因此 pgNumType 会被继承下来；必须先清掉再按需重设，否则后续节
        会把页码重新从 1 开始（目录/正文页码错乱）。
        """
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.8)

        # 清除继承来的页码设置
        for old in sec._sectPr.findall(qn("w:pgNumType")):
            sec._sectPr.remove(old)

        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.text = ""
        _fmt(hp, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        if header:
            r = hp.add_run(HEADER_TEXT)
            set_font(r, "宋体", 9)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        _fmt(fp, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        if page_no:
            add_field(fp, "PAGE")
        if page_start is not None:
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), str(page_start))
            sec._sectPr.append(pg)

    def new_section(self, header=True, page_start=None, page_no=False):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.config_section(sec, header, page_start, page_no)
        return sec

    # ---------- 段落
    def p(self, text, size=12, indent=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5,
          before=0, after=0, bold=False, cn="宋体"):
        par = self.doc.add_paragraph()
        _fmt(par, line=line, before=before, after=after, indent=indent, align=align, size=size)
        r = par.add_run(text)
        set_font(r, cn, size, bold)
        return par

    def h1(self, text, page_break=False):
        if page_break:
            self.doc.add_page_break()
        par = self.doc.add_paragraph()
        _fmt(par, line=1.5, before=12, after=18, align=WD_ALIGN_PARAGRAPH.CENTER, size=16)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(text)
        set_font(r, "黑体", 16, True)
        return par

    def h2(self, text):
        par = self.doc.add_paragraph()
        _fmt(par, line=1.5, before=10, after=12, align=WD_ALIGN_PARAGRAPH.LEFT, size=14)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(text)
        set_font(r, "黑体", 14, True)
        return par

    def h3(self, text):
        par = self.doc.add_paragraph()
        _fmt(par, line=1.5, before=6, after=6, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(text)
        set_font(r, "黑体", 12, True)
        return par

    def blank(self, size=12):
        par = self.doc.add_paragraph()
        _fmt(par, line=1.0, size=size)
        return par

    # ---------- 图
    def fig(self, filename, caption, width=14.6, source=None):
        # 优先取 figures/report/ 下的“裸图”（无图内标题、无图内来源）
        path = os.path.join(FIG, "report", filename)
        if not os.path.exists(path):
            path = os.path.join(FIG, filename)
        par = self.doc.add_paragraph()
        _fmt(par, line=1.0, before=6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        par.paragraph_format.keep_with_next = True      # 图、图题、来源三者不跨页拆开
        par.add_run().add_picture(path, width=Cm(width))

        cap = self.doc.add_paragraph()
        _fmt(cap, line=1.0, before=2, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        cap.paragraph_format.keep_with_next = True
        r = cap.add_run(caption)
        set_font(r, "宋体", 9, True)

        if source:
            s = self.doc.add_paragraph()
            _fmt(s, line=1.0, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5)
            rs = s.add_run(source)
            set_font(rs, "宋体", 7.5)

    # ---------- 表
    def tbl(self, df, caption, note=None, col_widths=None, size=9, first_left=True):
        cap = self.doc.add_paragraph()
        _fmt(cap, line=1.0, before=8, after=3, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        cap.paragraph_format.keep_with_next = True      # 表题与表体不跨页拆开
        r = cap.add_run(caption)
        set_font(r, "宋体", 9, True)

        nrow, ncol = df.shape[0] + 1, df.shape[1]
        table = self.doc.add_table(rows=nrow, cols=ncol)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_borders(table, sz=4)

        widths = col_widths or [16.5 / ncol] * ncol
        assert sum(widths) <= 16.8, f"{caption} 列宽合计 {sum(widths):.2f}cm 超出 17cm 版心"
        for j, w in enumerate(widths):
            for cell in table.columns[j].cells:
                cell.width = Cm(w)
        fix_table_width(table, widths)

        for j, c in enumerate(df.columns):
            cell_text(table.rows[0].cells[j], c, size=size, bold=True)
        repeat_header(table.rows[0])
        for i in range(df.shape[0]):
            for j in range(ncol):
                v = df.iat[i, j]
                if isinstance(v, float):
                    v = f"{v:,.2f}"
                al = WD_ALIGN_PARAGRAPH.LEFT if (first_left and j == 0) else WD_ALIGN_PARAGRAPH.CENTER
                cell_text(table.rows[i + 1].cells[j], v, size=size, align=al)

        if note:
            s = self.doc.add_paragraph()
            _fmt(s, line=1.0, before=2, after=8, align=WD_ALIGN_PARAGRAPH.LEFT, size=7.5)
            rs = s.add_run(note)
            set_font(rs, "宋体", 7.5)
        else:
            self.blank(6)
        return table

    def save(self, path):
        self._merge_section_breaks()
        self.doc.save(path)
        return path

    def _merge_section_breaks(self):
        """python-docx 的分节符寄生在一个独立空段落里。该段落按正文行高占位，
        若前一节恰好排满整页，它会被挤到下一页，凭空多出一整页空白。
        这里把 sectPr 搬到前一个正文段落的 pPr 末尾并删掉这个空段落；
        若搬不动（前一个不是段落），则把它压成 1pt。"""
        body = self.doc.element.body
        for p in list(body.findall(qn("w:p"))):
            pPr = p.find(qn("w:pPr"))
            if pPr is None:
                continue
            sectPr = pPr.find(qn("w:sectPr"))
            if sectPr is None:
                continue
            if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip():
                continue          # 有文字的分节段落不能动
            prev = p.getprevious()
            if prev is not None and prev.tag == qn("w:p"):
                prev_pPr = prev.find(qn("w:pPr"))
                if prev_pPr is None:
                    prev_pPr = OxmlElement("w:pPr")
                    prev.insert(0, prev_pPr)
                pPr.remove(sectPr)
                prev_pPr.append(sectPr)
                body.remove(p)
            else:
                from docx.text.paragraph import Paragraph
                pf = Paragraph(p, self.doc._body).paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = Pt(1)


# ------------------------------------------------------------------ 任务书（取自模板，保证与任务原文一致）
def read_taskbook():
    tpl = Document(TPL)
    cell = tpl.tables[0].rows[0].cells[0]
    lines = [p.text for p in cell.paragraphs]
    return lines
